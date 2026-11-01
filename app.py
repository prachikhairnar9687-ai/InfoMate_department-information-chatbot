import csv
import io
import os
import re
import time
import uuid
from datetime import timedelta
from functools import wraps
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from flask import (
    Flask, render_template, request, jsonify, redirect,
    url_for, session, flash
)
from supabase import create_client
from google import genai
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook

# OCR is optional. The app can still run without pytesseract/Tesseract installed.
try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env", override=True)

app = Flask(__name__)
app.secret_key = os.getenv(
    "FLASK_SECRET_KEY", "CHANGE_THIS_SECRET_BEFORE_DEPLOYING"
)
app.permanent_session_lifetime = timedelta(hours=8)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024

SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "").strip()
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY", "").strip()
BUCKET = os.getenv("SUPABASE_BUCKET", "department-documents").strip()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite").strip()

supabase = (
    create_client(SUPABASE_URL, SUPABASE_SERVICE_KEY)
    if SUPABASE_URL and SUPABASE_SERVICE_KEY
    else None
)
gemini = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

ALLOWED_EXTENSIONS = {
    "pdf", "docx", "txt", "csv", "xlsx", "xls",
    "png", "jpg", "jpeg"
}

# Keep text extraction bounded to protect memory and prompt preparation.
# For very large datasets, store/query the dataset in a database instead.
MAX_EXTRACTED_CHARS = 1_500_000
CACHE = {}
CACHE_SECONDS = 60

STOP_WORDS = {
    "the", "is", "are", "was", "were", "a", "an", "and", "or", "of",
    "to", "in", "on", "for", "me", "my", "we", "our", "you", "your",
    "what", "when", "where", "how", "can", "could", "would", "please",
    "tell", "show", "give", "about", "from", "with", "this", "that",
    "it", "do", "does", "did", "i", "us", "have", "has", "had",
    "find", "get", "want", "need", "department", "information"
}


# ---------------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------------

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("admin_user"):
            return redirect(url_for("admin_login"))
        return fn(*args, **kwargs)
    return wrapper


def sanitize_text(text):
    """Replace invalid Unicode surrogate characters."""
    return str(text).encode("utf-8", errors="replace").decode("utf-8")


def file_extension(filename):
    return Path(filename).suffix.lower().lstrip(".")


def tokenize(text):
    return {
        token for token in re.findall(r"[a-zA-Z0-9]{2,}", (text or "").lower())
        if token not in STOP_WORDS
    }


def extract_text(filename, raw):
    """Extract searchable text from supported uploaded files."""
    ext = file_extension(filename)

    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(raw))
            parts = []
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    parts.append(f"PAGE {page_number}\n{page_text}")
            text = "\n\n".join(parts)

        elif ext == "docx":
            doc = DocxDocument(io.BytesIO(raw))
            parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            # Include tables, which commonly contain schedules and academic data.
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)

        elif ext == "txt":
            text = raw.decode("utf-8-sig", errors="replace")

        elif ext == "csv":
            text = raw.decode("utf-8-sig", errors="replace")
            # Validate that it can be parsed as a CSV, but preserve original rows.
            try:
                list(csv.reader(io.StringIO(text)))
            except csv.Error as exc:
                raise ValueError(f"Invalid CSV file: {exc}") from exc

        elif ext == "xlsx":
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=True
            )
            parts = []
            try:
                for sheet in workbook.worksheets:
                    parts.append(f"SHEET: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = [
                            "" if value is None else str(value)
                            for value in row
                        ]
                        if any(value.strip() for value in values):
                            parts.append(" | ".join(values))
            finally:
                workbook.close()
            text = "\n".join(parts)

        elif ext == "xls":
            # The xlrd package is required for legacy .xls workbooks.
            workbook = pd.ExcelFile(io.BytesIO(raw), engine="xlrd")
            parts = []
            for sheet_name in workbook.sheet_names:
                frame = pd.read_excel(workbook, sheet_name=sheet_name)
                parts.append(f"SHEET: {sheet_name}")
                parts.append(frame.to_csv(index=False))
            text = "\n".join(parts)

        elif ext in {"png", "jpg", "jpeg"}:
            if Image is None or pytesseract is None:
                raise ValueError(
                    "Image OCR is not configured. Install Pillow and pytesseract, "
                    "and install the Tesseract OCR application."
                )
            image_file = Image.open(io.BytesIO(raw))
            text = pytesseract.image_to_string(image_file)

        else:
            raise ValueError("Unsupported file format.")

        return sanitize_text(text)[:MAX_EXTRACTED_CHARS]

    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(f"Could not read this file: {exc}") from exc


def get_documents(published_only=True):
    """Fetch document rows in pages instead of silently relying on one API page."""
    if not supabase:
        return []

    results = []
    page_size = 500
    start = 0

    while True:
        query = (
            supabase.table("documents")
            .select("id,name,category,content,created_at,status,storage_path")
            .order("created_at", desc=True)
            .range(start, start + page_size - 1)
        )
        if published_only:
            query = query.eq("status", "published")

        batch = query.execute().data or []
        results.extend(batch)

        if len(batch) < page_size:
            break

        start += page_size

    return results


def get_notices(published_only=True):
    """Fetch notices in pages, newest first."""
    if not supabase:
        return []

    results = []
    page_size = 500
    start = 0

    while True:
        query = (
            supabase.table("notices")
            .select("id,title,body,category,published,created_at")
            .order("created_at", desc=True)
            .range(start, start + page_size - 1)
        )
        if published_only:
            query = query.eq("published", True)

        batch = query.execute().data or []
        results.extend(batch)

        if len(batch) < page_size:
            break

        start += page_size

    return results


# ---------------------------------------------------------------------------
# Retrieval: published notices and documents
# ---------------------------------------------------------------------------

def retrieve_context(question, limit=8):
    """Retrieve relevant published notices and documents."""

    docs = get_documents(published_only=True)
    notices = get_notices(published_only=True)

    q = (question or "").lower().strip()

    # Recognize questions asking for the latest or recent notices.
    latest_notice_patterns = [
        r"\b(latest|newest|most recent|recent|new|current|latest available)\b.{0,60}\b(notice|notices|announcement|announcements|circular|circulars|update|updates)\b",
        r"\b(notice|notices|announcement|announcements|circular|circulars|update|updates)\b.{0,60}\b(latest|newest|most recent|recent|new|current)\b",
        r"\b(any|what|show|tell|give|list|display)\b.{0,50}\b(new|latest|recent)\b",
        r"\bwhat\s+is\s+the\s+latest\s+notice\b",
        r"\bany\s+(new|latest|recent)\s+(department\s+)?(notice|update|announcement|circular)s?\b",
    ]

    asks_for_latest_notice = any(
        re.search(pattern, q)
        for pattern in latest_notice_patterns
    )

    # IMPORTANT: For latest-notice questions, do not use keyword matching.
    # The notices are already ordered newest first by get_notices().
    if asks_for_latest_notice:
        if not notices:
            return (
                "No published notices are currently available in the "
                "department notice database.",
                False
            )

        selected_notices = notices[:3]

        notice_context = []

        for index, notice in enumerate(selected_notices, start=1):
            notice_context.append(
                f"SOURCE TYPE: PUBLISHED DEPARTMENT NOTICE\n"
                f"NOTICE RANK: {index}\n"
                f"TITLE: {notice.get('title') or 'Untitled notice'}\n"
                f"CATEGORY: {notice.get('category') or 'General'}\n"
                f"PUBLISHED DATE: {notice.get('created_at') or 'Date not available'}\n"
                f"NOTICE CONTENT:\n{notice.get('body') or 'No notice details available.'}"
            )

        return "\n\n---\n\n".join(notice_context), True

    # Normal keyword-based retrieval for other questions.
    terms = tokenize(question)

    if not terms:
        return "", False

    ranked = []

    for doc in docs:
        name = doc.get("name") or ""
        category = doc.get("category") or ""
        content = doc.get("content") or ""

        metadata = f"{name} {category}"
        searchable = f"{metadata}\n{content}".lower()

        matches = {
            term for term in terms
            if term in searchable
        }

        if not matches:
            continue

        metadata_lower = metadata.lower()

        score = len(matches)
        score += 3 * sum(
            1 for term in terms
            if term in metadata_lower
        )

        chunks = re.split(r"(?<=[.!?])\s+|\n+", content)

        matching_chunks = [
            chunk.strip()
            for chunk in chunks
            if chunk.strip()
            and any(term in chunk.lower() for term in terms)
        ]

        excerpt = "\n".join(matching_chunks[:50]) or content[:3000]

        ranked.append((
            score,
            f"SOURCE TYPE: DOCUMENT\n"
            f"NAME: {name}\n"
            f"CATEGORY: {category}\n"
            f"CONTENT:\n{excerpt}"
        ))

    for notice in notices:
        title = notice.get("title") or ""
        body = notice.get("body") or ""
        category = notice.get("category") or ""

        searchable = f"{title}\n{body}\n{category}".lower()

        matches = {
            term for term in terms
            if term in searchable
        }

        if not matches:
            continue

        score = len(matches)

        score += 4 * sum(
            1 for term in terms
            if term in title.lower()
        )

        ranked.append((
            score,
            f"SOURCE TYPE: PUBLISHED NOTICE\n"
            f"TITLE: {title}\n"
            f"CATEGORY: {category}\n"
            f"PUBLISHED DATE: {notice.get('created_at') or 'Date not available'}\n"
            f"NOTICE CONTENT:\n{body}"
        ))

    ranked.sort(key=lambda item: item[0], reverse=True)

    selected = ranked[:limit]

    context = "\n\n---\n\n".join(
        item[1] for item in selected
    )

    return context, bool(selected)


def wants_document(question):
    """Detect an explicit request for a file/link, rather than every question."""
    q = (question or "").lower()

    patterns = [
        r"\b(show|send|give|share|provide|attach|open|download|get)\b.{0,70}"
        r"\b(pdf|document|file|notes|sheet|spreadsheet|excel|csv|syllabus|"
        r"timetable|circular|notice)\b",
        r"\b(pdf|document|file|notes|sheet|spreadsheet|excel|csv|syllabus|"
        r"timetable)\b.{0,70}\b(send|share|download|open|attach|give|me)\b",
        r"\b(where can i|how can i)\s+(find|download|open|get)\b",
        r"\b(need|want)\b.{0,50}\b(pdf|document|file|notes|spreadsheet|excel)\b",
    ]
    return any(re.search(pattern, q) for pattern in patterns)


def find_relevant_documents(question, limit=3):
    """Return published documents only when a document request is detected."""
    docs = get_documents(published_only=True)
    terms = tokenize(question)
    if not terms:
        return []

    ranked = []
    for doc in docs:
        name = doc.get("name") or ""
        category = doc.get("category") or ""
        content = doc.get("content") or ""
        metadata = f"{name} {category}".lower()
        searchable = f"{metadata}\n{content}".lower()

        matched = {term for term in terms if term in searchable}
        if not matched:
            continue

        score = len(matched)
        score += 4 * sum(1 for term in terms if term in metadata)

        # Require at least one match in filename/category, or at least
        # two distinct content matches, to reduce irrelevant file links.
        metadata_matches = sum(1 for term in terms if term in metadata)
        if metadata_matches == 0 and len(matched) < 2:
            continue

        ranked.append({
            "id": doc.get("id"),
            "name": name,
            "category": category or "General",
            "score": score,
        })

    ranked.sort(key=lambda item: item["score"], reverse=True)
    return [
        {"id": item["id"], "name": item["name"], "category": item["category"]}
        for item in ranked[:limit]
    ]


# ---------------------------------------------------------------------------
# Structured CSV/Excel analysis
# ---------------------------------------------------------------------------

def _read_structured_document(doc):
    """Load a CSV or Excel file from its indexed text representation."""
    name = (doc.get("name") or "").lower()
    content = doc.get("content") or ""
    if not content.strip():
        return None

    try:
        if name.endswith(".csv"):
            return {"data": pd.read_csv(io.StringIO(content)), "sheet": None}

        if name.endswith(".xlsx"):
            # Extracted workbook text is represented as SHEET: ... and
            # pipe-separated rows. Parse each sheet block conservatively.
            blocks = re.split(r"(?m)^SHEET:\s*", content)
            parsed = []
            for block in blocks:
                lines = [line for line in block.splitlines() if line.strip()]
                if not lines:
                    continue
                sheet_name = lines[0].strip()
                rows = []
                for line in lines[1:]:
                    rows.append([cell.strip() for cell in line.split("|")])
                if rows:
                    header = rows[0]
                    parsed.append((sheet_name, pd.DataFrame(rows[1:], columns=header)))
            return {"sheets": parsed, "data": parsed[0][1] if parsed else pd.DataFrame(), "sheet": None}

        if name.endswith(".xls"):
            blocks = re.split(r"(?m)^SHEET:\s*", content)
            parsed = []
            for block in blocks:
                lines = [line for line in block.splitlines() if line.strip()]
                if not lines:
                    continue
                sheet_name = lines[0].strip()
                try:
                    frame = pd.read_csv(io.StringIO("\n".join(lines[1:])))
                except Exception:
                    continue
                parsed.append((sheet_name, frame))
            return {"sheets": parsed, "data": parsed[0][1] if parsed else pd.DataFrame(), "sheet": None}

    except Exception:
        app.logger.exception("Could not parse structured file %s", name)

    return None


def answer_from_structured_data(question):
    """
    Basic full-dataset analytics for CSV/Excel files.
    Supports record counts and simple numeric min/max/sum/average queries.
    It deliberately returns None for unsupported queries so retrieval/Gemini
    can handle the question rather than inventing a result.
    """
    q = (question or "").lower()
    is_count = bool(re.search(r"\b(how many|count|number of|total number of|records|rows)\b", q))
    operation = None

    if re.search(r"\b(average|mean|avg)\b", q):
        operation = "average"
    elif re.search(r"\b(sum|total|combined)\b", q):
        operation = "sum"
    elif re.search(r"\b(highest|maximum|max|largest|most)\b", q):
        operation = "max"
    elif re.search(r"\b(lowest|minimum|min|smallest|least)\b", q):
        operation = "min"

    if not is_count and operation is None:
        return None

    docs = get_documents(published_only=True)
    question_terms = tokenize(q)
    candidates = []

    for doc in docs:
        filename = doc.get("name") or ""
        if not filename.lower().endswith((".csv", ".xlsx", ".xls")):
            continue

        parsed = _read_structured_document(doc)
        if not parsed:
            continue

        frames = parsed.get("sheets") or [(None, parsed.get("data"))]
        for sheet_name, frame in frames:
            if frame is None or frame.empty:
                continue

            metadata_terms = tokenize(
                f"{filename} {doc.get('category', '')} {sheet_name or ''} "
                + " ".join(str(col) for col in frame.columns)
            )
            overlap = len(question_terms.intersection(metadata_terms))

            # If there are several structured files, prefer a dataset whose
            # name, category, sheet name, or headers match the question.
            candidates.append((overlap, doc, sheet_name, frame))

    if not candidates:
        return None

    candidates.sort(key=lambda item: item[0], reverse=True)
    best_overlap = candidates[0][0]
    if len(candidates) > 1 and best_overlap == 0:
        return None

    _, doc, sheet_name, frame = candidates[0]
    filename = doc.get("name") or "dataset"

    if is_count:
        # Count all data rows in the complete stored dataset.
        return (
            f"According to the published dataset '{filename}', "
            f"there are {len(frame)} records"
            + (f" in sheet '{sheet_name}'." if sheet_name else ".")
        )

    # Find the requested numeric column by matching question words.
    numeric = frame.copy()
    for column in numeric.columns:
        numeric[column] = pd.to_numeric(numeric[column], errors="ignore")

    numeric_columns = [
        column for column in numeric.columns
        if pd.api.types.is_numeric_dtype(numeric[column])
    ]

    requested_columns = [
        column for column in numeric_columns
        if str(column).lower() in q
        or any(token in str(column).lower() for token in question_terms)
    ]

    if requested_columns:
        column = requested_columns[0]
    elif len(numeric_columns) == 1:
        column = numeric_columns[0]
    else:
        return None

    values = pd.to_numeric(numeric[column], errors="coerce").dropna()
    if values.empty:
        return None

    if operation == "average":
        result = values.mean()
        label = "Average"
    elif operation == "sum":
        result = values.sum()
        label = "Total"
    elif operation == "max":
        result = values.max()
        label = "Maximum"
    elif operation == "min":
        result = values.min()
        label = "Minimum"
    else:
        return None

    suffix = f" in sheet '{sheet_name}'" if sheet_name else ""
    return (
        f"{label} of '{column}' in the published dataset '{filename}'"
        f"{suffix}: {result:g} (calculated from {len(values)} numeric values)."
    )


# ---------------------------------------------------------------------------
# Answer generation
# ---------------------------------------------------------------------------

def deterministic_answer(question, context, has_context):
    q = question.lower().strip()

    if re.fullmatch(
        r"(hi|hello|hey|hii|good morning|good afternoon|good evening)[!. ]*",
        q
    ):
        return (
            "Hello! 😊 I'm InfoMate, your Computer Department academic assistant. "
            "You can ask me about timetables, exam dates, notices, syllabi, "
            "and academic schedules."
        )

    if any(phrase in q for phrase in (
        "what can you do", "help me", "what can i ask"
    )):
        return (
            "I can help you find department information such as timetables, "
            "examination dates, academic calendars, notices, syllabi, and "
            "course details. Ask a specific question or request a document."
        )

    if not has_context:
        return (
            "I couldn't find this information in the currently published "
            "department documents or notices. Please try different keywords "
            "or check with the department office. I won't guess academic "
            "dates or room numbers."
        )

    return None


def answer_question(question):
    key = question.strip().lower()
    now = time.time()

    if key in CACHE and now - CACHE[key][0] < CACHE_SECONDS:
        return CACHE[key][1], "cache"

    # Structured data analysis happens before text retrieval.
    try:
        structured_answer = answer_from_structured_data(question)
        if structured_answer:
            CACHE[key] = (now, structured_answer)
            return structured_answer, "data_analysis"
    except Exception:
        app.logger.exception("Structured data analysis failed")

    context, found = retrieve_context(question)
    fallback = deterministic_answer(question, context, found)

    if fallback:
        CACHE[key] = (now, fallback)
        return fallback, "local"

    if not gemini:
        if found and context:
            return (
                "The AI service is not configured. Here is the relevant "
                "published source information:\n\n" + context[:5000],
                "retrieval"
            )
        return (
            "I couldn't find a relevant answer in the published department "
            "documents or notices. Please try different keywords or check "
            "with the department office.",
            "no_source"
        )

    prompt = f"""You are InfoMate, a Computer Department academic assistant.

Answer using ONLY the approved department sources below. Never invent dates,
seat numbers, rooms, schedules, or policies. If the sources do not answer the
question, clearly say that the information is not available in the published
sources. Distinguish notices from documents. Use the latest notice only when
the question asks for the latest/current notice. Be concise and friendly.

QUESTION:
{question}

APPROVED PUBLISHED DEPARTMENT SOURCES:
{context[:18000]}

If no source content is relevant enough to answer, say so instead of guessing.
"""

    try:
        response = gemini.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt
        )
        result = (response.text or "").strip()
        if not result:
            result = "I couldn't generate a response just now. Please try again."

        CACHE[key] = (now, result)
        return result, "gemini"

    except Exception:
        app.logger.exception("Gemini generation failed")

        if found and context.strip():
            result = (
                "The AI service is temporarily unavailable. Here is the "
                "relevant information from published department sources:\n\n"
                + context[:5000]
                + "\n\nPlease verify important dates and instructions against "
                "the official source."
            )
            CACHE[key] = (now, result)
            return result, "source_fallback"

        return (
            "I couldn't find a relevant answer in the published department "
            "documents or notices. Please try different keywords or check "
            "with the department office.",
            "no_source"
        )


# ---------------------------------------------------------------------------
# Student pages and APIs
# ---------------------------------------------------------------------------

@app.get("/")
def home():
    try:
        notices = get_notices(True)[:5]
    except Exception:
        app.logger.exception("Could not load homepage notices")
        notices = []
    return render_template("index.html", notices=notices)


@app.get("/chat")
def chat_page():
    return render_template("chat.html")


@app.post("/api/chat")
def chat_api():
    data = request.get_json(silent=True) or {}
    question = (data.get("message") or "").strip()

    if not question:
        return jsonify({
            "answer": "Please type a question first.",
            "source": "local",
            "documents": []
        }), 400

    if len(question) > 1200:
        return jsonify({
            "answer": "Please keep your question under 1,200 characters.",
            "source": "local",
            "documents": []
        }), 400

    try:
        answer, source = answer_question(question)

        documents = []
        if wants_document(question):
            try:
                documents = find_relevant_documents(question)
            except Exception:
                app.logger.exception("Relevant document lookup failed")
                documents = []

        return jsonify({
            "answer": answer,
            "source": source,
            "documents": documents
        })

    except Exception:
        app.logger.exception("Chat request failed")
        return jsonify({
            "answer": "I couldn't access department information right now. Please try again later.",
            "source": "fallback",
            "documents": []
        }), 500


@app.get("/api/documents/<document_id>/open")
def open_published_document(document_id):
    """Open any published file using a short-lived signed storage URL."""
    if not supabase:
        return jsonify({"error": "Document storage is unavailable."}), 503

    try:
        result = (
            supabase.table("documents")
            .select("id,name,status,storage_path")
            .eq("id", document_id)
            .eq("status", "published")
            .single()
            .execute()
        )
        doc = result.data

        if not doc:
            return jsonify({"error": "Document not found."}), 404

        storage_path = doc.get("storage_path") or ""
        if not storage_path:
            return jsonify({"error": "Document file is unavailable."}), 404

        signed = supabase.storage.from_(BUCKET).create_signed_url(
            storage_path, 300
        )
        url = signed.get("signedURL") or signed.get("signedUrl")

        if not url:
            return jsonify({"error": "Could not create a document link."}), 500

        return redirect(url)

    except Exception:
        app.logger.exception("Published document opening failed")
        return jsonify({"error": "Could not open this document right now."}), 500


@app.get("/health")
def health():
    return jsonify({
        "status": "ok",
        "supabase_configured": bool(supabase),
        "gemini_configured": bool(gemini),
        "gemini_model": GEMINI_MODEL,
        "ocr_configured": bool(Image and pytesseract),
        "note": "Configuration presence only; this does not verify provider authentication."
    })


# ---------------------------------------------------------------------------
# Admin authentication
# ---------------------------------------------------------------------------

@app.route("/admin", methods=["GET"])
def admin_login():
    if session.get("admin_user"):
        return redirect(url_for("admin_dashboard"))
    return render_template("admin_login.html")


@app.post("/admin/login")
def admin_login_post():
    email = (request.form.get("email") or "").strip()
    password = request.form.get("password") or ""

    if not SUPABASE_URL or not SUPABASE_ANON_KEY:
        flash(
            "Supabase Auth is not configured. Check server environment variables.",
            "error"
        )
        return redirect(url_for("admin_login"))

    try:
        auth_client = create_client(SUPABASE_URL, SUPABASE_ANON_KEY)
        result = auth_client.auth.sign_in_with_password({
            "email": email,
            "password": password
        })
        user = result.user

        if not user:
            raise ValueError("Sign-in did not return a user.")

        profile = (
            supabase.table("profiles")
            .select("role,active")
            .eq("id", user.id)
            .maybe_single()
            .execute()
            .data
        )

        allowed_roles = ("super_admin", "admin", "coordinator", "faculty")
        if (
            not profile
            or profile.get("active") is False
            or profile.get("role") not in allowed_roles
        ):
            auth_client.auth.sign_out()
            flash(
                "This account is not authorized to access the admin portal.",
                "error"
            )
            return redirect(url_for("admin_login"))

        session.clear()
        session.permanent = True
        session["admin_user"] = {
            "id": user.id,
            "email": user.email,
            "role": profile["role"]
        }
        return redirect(url_for("admin_dashboard"))

    except Exception:
        app.logger.exception("Admin sign-in failed")
        flash(
            "Sign-in failed. Check your email/password and confirm that the "
            "account has an authorized profile.",
            "error"
        )
        return redirect(url_for("admin_login"))


@app.post("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


@app.get("/admin/dashboard")
@admin_required
def admin_dashboard():
    try:
        docs = (
            supabase.table("documents")
            .select("id,name,category,status,created_at,storage_path")
            .order("created_at", desc=True)
            .execute().data or []
        )
        notices = (
            supabase.table("notices")
            .select("id,title,category,published,created_at")
            .order("created_at", desc=True)
            .execute().data or []
        )
        db_error = None
    except Exception as exc:
        app.logger.exception("Admin dashboard query failed")
        docs, notices, db_error = [], [], str(exc)

    return render_template(
        "admin_dashboard.html",
        docs=docs,
        notices=notices,
        db_error=db_error,
        admin=session["admin_user"],
        ai_ready=bool(gemini)
    )


# ---------------------------------------------------------------------------
# Admin document upload, open, delete
# ---------------------------------------------------------------------------

@app.post("/admin/documents")
@admin_required
def upload_document():
    file = request.files.get("file")
    category = (request.form.get("category") or "General").strip()[:80]

    if not file or not file.filename:
        flash("Choose a file first.", "error")
        return redirect(url_for("admin_dashboard"))

    if not allowed_file(file.filename):
        flash(
            "Allowed formats: PDF, DOCX, TXT, CSV, XLSX, XLS, PNG, JPG, JPEG.",
            "error"
        )
        return redirect(url_for("admin_dashboard"))

    filename = os.path.basename(file.filename).replace("/", "_").replace("\\", "_")
    raw = file.read()

    if not raw:
        flash("The selected file is empty.", "error")
        return redirect(url_for("admin_dashboard"))

    try:
        content = sanitize_text(extract_text(filename, raw))
        if not content.strip():
            flash(
                "No readable text was found. Images may need OCR, and scanned "
                "PDFs may need OCR before upload.",
                "error"
            )
            return redirect(url_for("admin_dashboard"))

        if not supabase:
            flash("Supabase is not configured on the server.", "error")
            return redirect(url_for("admin_dashboard"))

        document_id = str(uuid.uuid4())
        storage_path = f"{document_id}/{filename}"

        supabase.storage.from_(BUCKET).upload(
            storage_path,
            raw,
            {
                "content-type": file.mimetype or "application/octet-stream",
                "upsert": "false"
            }
        )

        supabase.table("documents").insert({
            "id": document_id,
            "name": filename,
            "category": category,
            "content": content,
            "status": "published",
            "storage_path": storage_path,
            "uploaded_by": session["admin_user"]["id"]
        }).execute()

        CACHE.clear()
        flash(f"{filename} uploaded and indexed successfully.", "success")

    except Exception:
        app.logger.exception("Document upload failed")
        flash(
            "Upload failed. Check the Supabase bucket, documents table, file "
            "size, required Python packages, and server logs.",
            "error"
        )

    return redirect(url_for("admin_dashboard"))


@app.get("/admin/documents/<document_id>/open")
@admin_required
def open_document(document_id):
    try:
        row = (
            supabase.table("documents")
            .select("storage_path,name")
            .eq("id", document_id)
            .single()
            .execute().data
        )
        signed = supabase.storage.from_(BUCKET).create_signed_url(
            row["storage_path"], 300
        )
        url = signed.get("signedURL") or signed.get("signedUrl")

        if not url:
            flash("Could not create a temporary file link.", "error")
            return redirect(url_for("admin_dashboard"))

        return redirect(url)

    except Exception:
        app.logger.exception("Admin document opening failed")
        flash(
            "Could not open the file. Check that the storage object exists.",
            "error"
        )
        return redirect(url_for("admin_dashboard"))


@app.post("/admin/documents/<document_id>/delete")
@admin_required
def delete_document(document_id):
    try:
        row = (
            supabase.table("documents")
            .select("storage_path")
            .eq("id", document_id)
            .single()
            .execute().data
        )

        if row and row.get("storage_path"):
            try:
                supabase.storage.from_(BUCKET).remove([row["storage_path"]])
            except Exception:
                app.logger.warning(
                    "Storage deletion failed; continuing with database deletion."
                )

        supabase.table("documents").delete().eq("id", document_id).execute()
        CACHE.clear()
        flash("Document deleted.", "success")

    except Exception:
        app.logger.exception("Document deletion failed")
        flash(
            "Could not delete this document. Check Supabase permissions and logs.",
            "error"
        )

    return redirect(url_for("admin_dashboard"))


# ---------------------------------------------------------------------------
# Admin notices
# ---------------------------------------------------------------------------

@app.post("/admin/notices")
@admin_required
def add_notice():
    title = (request.form.get("title") or "").strip()[:180]
    body = (request.form.get("body") or "").strip()[:12000]
    category = (request.form.get("category") or "General").strip()[:80]
    published = request.form.get("published") == "on"

    if not title or not body:
        flash("Notice title and content are required.", "error")
        return redirect(url_for("admin_dashboard"))

    try:
        supabase.table("notices").insert({
            "title": title,
            "body": body,
            "category": category,
            "published": published,
            "created_by": session["admin_user"]["id"]
        }).execute()

        CACHE.clear()
        flash("Notice saved.", "success")

    except Exception:
        app.logger.exception("Notice creation failed")
        flash("Could not save notice. Check the database migration.", "error")

    return redirect(url_for("admin_dashboard"))


@app.post("/admin/notices/<notice_id>/delete")
@admin_required
def delete_notice(notice_id):
    try:
        supabase.table("notices").delete().eq("id", notice_id).execute()
        CACHE.clear()
        flash("Notice deleted.", "success")
    except Exception:
        app.logger.exception("Notice deletion failed")
        flash("Could not delete notice.", "error")

    return redirect(url_for("admin_dashboard"))


if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.getenv("PORT", "5000")),
        debug=os.getenv("FLASK_DEBUG", "false").lower() == "true"
    )
